"""Stateless API and bound frontend regression checks, without live collection."""
import base64
import csv
import io

import pytest
from openpyxl import load_workbook

import app as api
from gui import vercel_app as gui


def test_stateless_demo_has_complete_results_and_downloads():
    response = api.app.test_client().post(
        "/run", json={"vertical": "web_design", "demo": True})
    assert response.status_code == 200
    body = response.get_json()
    assert body["done"] and body["error"] is None and body["stats"]["total"] == 5
    assert "job_id" not in body
    csv_bytes = base64.b64decode(body["exports"]["csv"]["data"])
    rows = list(csv.DictReader(io.StringIO(csv_bytes.decode("utf-8-sig"))))
    assert len(rows) == 5 and any(r["Business"] == "Summit Plumbing Co" for r in rows)
    workbook = load_workbook(io.BytesIO(base64.b64decode(body["exports"]["xlsx"]["data"])))
    assert workbook.active.max_row == 6
    workbook.close()
    assert response.headers["Cache-Control"] == "no-store"


@pytest.mark.parametrize("params", [
    [], {"vertical": "not_a_vertical", "demo": True},
    {"vertical": "au_video_ops"},
    {"vertical": "au_video_ops_overture", "market": "Sydney"},
    {"demo": True, "sources": ["overture"]},
    {"demo": "true"}, {"demo": True, "limit": -1},
    {"demo": True, "enrich_cap": "bad"},
])
def test_invalid_requests_do_not_run_collection(params, monkeypatch):
    monkeypatch.setattr(api, "run_pipeline", lambda *a, **k: pytest.fail("collection attempted"))
    assert api.app.test_client().post("/run", json=params).status_code == 400


def test_hosted_choices_exclude_overture_only_vertical():
    keys = {v["key"] for v in api.app.test_client().get("/metadata").get_json()["verticals"]}
    assert "au_video_ops" in keys and "au_video_ops_overture" not in keys


def test_failed_osm_source_is_a_failed_request(monkeypatch):
    def failed(*args, **kwargs):
        raise RuntimeError("OSM collection failed")

    monkeypatch.setattr(api, "run_pipeline", failed)
    response = api.app.test_client().post("/run", json={"market": "Sydney"})
    assert response.status_code == 502
    assert "OSM collection failed" in response.get_json()["error"]


def test_response_and_exports_are_bounded(monkeypatch):
    monkeypatch.setattr(api, "run_pipeline", lambda *a, **k: [
        {"name": f"Video Co {i}", "tier": "C", "score": 10} for i in range(250)])
    body = api.app.test_client().post("/run", json={"market": "Sydney"}).get_json()
    assert body["stats"]["total"] == 200 and len(body["leads"]) == 200
    rows = list(csv.reader(io.StringIO(base64.b64decode(
        body["exports"]["csv"]["data"]).decode())))
    assert len(rows) == 201


def test_missing_binding_has_no_localhost_fallback(monkeypatch):
    monkeypatch.delenv("LEADGEN_API_URL", raising=False)
    response = gui.app.test_client().get("/health")
    assert response.status_code == 503
    assert "binding is missing" in response.get_json()["error"]


def test_runtime_binding_and_public_request_flow(monkeypatch):
    # This simulates Vercel injection for an isolated test; production configuration
    # never assigns a value to the bound variable.
    monkeypatch.setenv("LEADGEN_API_URL", "http://internal.test/service/app/")
    calls = []

    class Response:
        def __init__(self, response):
            self.response = response
            self.status_code = response.status_code

        def json(self):
            return self.response.get_json()

    def request(method, url, **kwargs):
        calls.append((method, url))
        assert kwargs["allow_redirects"] is False
        path = "/" + url.split("/service/app/", 1)[1]
        return Response(api.app.test_client().open(path, method=method, json=kwargs["json"]))

    monkeypatch.setattr(gui.requests, "request", request)
    client = gui.app.test_client()
    page = client.get("/")
    assert page.status_code == 200 and b"const SERVERLESS_MODE = true" in page.data
    assert b"au_video_ops" in page.data
    result = client.post("/run", json={"vertical": "web_design", "demo": True})
    assert result.status_code == 200 and result.get_json()["stats"]["total"] == 5
    assert calls == [("GET", "http://internal.test/service/app/metadata"),
                     ("POST", "http://internal.test/service/app/run")]
    assert client.get("/metadata").status_code == 404
    assert client.get("/recent").get_json() == []


def test_api_errors_are_preserved_by_gui(monkeypatch):
    monkeypatch.setattr(gui, "call_api", lambda *a, **k: ({"error": "Bad market"}, 400))
    response = gui.app.test_client().post("/run", json={"market": "bad"})
    assert response.status_code == 400 and response.get_json()["error"] == "Bad market"
